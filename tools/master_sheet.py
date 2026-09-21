#!/usr/bin/env python3
"""Build `A7 Master Sheet 2026-27.xlsx`: the year's plan with everything for each day in one place.

    python3 tools/master_sheet.py [out.xlsx]

Tabs: Today (the next class day, worked out from today's date, with tap-to-open links) ·
Scope & Sequence (one row per class day: links to every document, IXL, learning target, essential
question, vocabulary, benchmark wording, MTRs, closure question, Math Nation lessons) · IXL
tracker · Standards · Unit tests · Month calendar · Vocabulary · Units · Units 1–2 decks · About.

Reads, and never edits: the calendar table (a7/reference/A7 SCOPE AND SEQUENCE 2026-27.md, made by
scope_calendar.py), the IXL due-date sheet, the Source of Truth for benchmark wording, the Math
Nation table of contents, the lesson and unit specs in a7/build, and the package folders. Links
point at the GitHub repository. Rerun after any unit is built; the file records its commit.
"""
import os, re, sys, glob, datetime, subprocess, urllib.parse, importlib.util, contextlib, io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.pagebreak import Break

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.normpath(os.path.join(HERE, ".."))
REF = os.path.join(ROOT, "a7", "reference")
BUILD = os.path.join(ROOT, "a7", "build")
PKG = os.path.join(ROOT, "a7", "packages")
CAL = os.path.join(REF, "A7 SCOPE AND SEQUENCE 2026-27.md")
DUE = os.path.join(REF, "A7 IXL DUE DATES 2026-27.md")
SOT = os.path.join(REF, "Florida BEST Grade 8 - Source of Truth.md")
TOC = os.path.join(REF, "MATH NATION A7 BOOK - table of contents and IXL plan.md")
REPO = "https://github.com/croix18/croix18-windy-hill-a7"
FONT = "Arial"
sys.path.insert(0, HERE)
sys.path.insert(0, BUILD)
import tex2text
with contextlib.redirect_stdout(io.StringIO()):
    import scope_calendar as SC            # the school days, the holidays, lesson -> book lessons
from lib import unitbuild                  # the unit-test ledger (question numbers and points)

UNIT_TITLES = {1: "Equations and Inequalities", 2: "Real Numbers, Square Roots and Cube Roots",
               3: "Exponents and Scientific Notation", 4: "Solving Problems with Rational Numbers",
               5: "Solving Multi-Step Problems with Proportional Relationships", 6: "Relationships in Triangles",
               7: "Representing Proportional Relationships", 8: "Relationships in Circles", 9: "Surface Area and Volume",
               10: "Linear Relationships", 11: "Solving Equations and Systems of Equations", 12: "Functions",
               13: "Representing Data", 14: "Equivalent Algebraic Expressions", 15: "Transformations",
               16: "Properties and Theorems of Angles", 17: "Probability"}
# Units 1 and 2 ran before the calendar starts (23 Sep); their benchmarks are the book's unit headings.
EARLY_UNITS = {1: ["7.NSO.1.2", "7.AR.2.2", "8.AR.2.2"], 2: ["8.AR.2.3", "8.NSO.1.1", "8.NSO.1.2"]}
# MA.K12.MTR titles, verbatim from Florida's B.E.S.T. Standards for Mathematics (pp. on the MTRs).
MTR = {"MTR.1.1": "Actively participate in effortful learning both individually and collectively.",
       "MTR.2.1": "Demonstrate understanding by representing problems in multiple ways.",
       "MTR.3.1": "Complete tasks with mathematical fluency.",
       "MTR.4.1": "Engage in discussions that reflect on the mathematical thinking of self and others.",
       "MTR.5.1": "Use patterns and structure to help understand and connect mathematical concepts.",
       "MTR.6.1": "Assess the reasonableness of solutions.",
       "MTR.7.1": "Apply mathematics to real-world contexts."}
MONTHS = {m: i for i, m in enumerate(["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"], 1)}
START = SC.start                           # 23 Sep 2026


# ---------------------------------------------------------------- sources
def date_of(s):
    """'Sep 23' or 'Wed Sep 23' -> a date in the 2026-27 school year."""
    mon, day = s.split()[-2:]
    m = MONTHS[mon]
    return datetime.date(2026 if m >= 7 else 2027, m, int(day))


def clean_md(s):
    s = re.sub(r"\*\*(.+?)\*\*", r"\1", s)
    s = re.sub(r"(?<![\w*])\*(?!\s)(.+?)(?<!\s)\*(?![\w*])", r"\1", s)
    return s.replace("`", "").strip()


def read_calendar():
    rows = []
    for ln in open(CAL, encoding="utf-8"):
        if not ln.startswith("| ") or ln.startswith("| Date"):
            continue
        date, day, unit, code, taught, bm, ixl, due = [c.strip() for c in ln.strip().strip("|").split("|")]
        rows.append(dict(date=date_of(date), day=day, unit=int(unit) if unit else None, code=code,
                         taught=clean_md(re.sub(r"^\*{1,2}\w+\*{1,2}\s*", "", taught)),
                         bm=bm, ixl=ixl, due=date_of(due) if due else None))
    return rows


def read_due_sheet():
    out = []
    for ln in open(DUE, encoding="utf-8"):
        if not ln.startswith("| ") or ln.startswith("| Assigned"):
            continue
        when, lessons, skills, due = [c.strip() for c in ln.strip().strip("|").split("|")]
        out.append(dict(assigned=date_of(when), lessons=lessons, skills=skills, due=date_of(due),
                        codes=re.findall(r"\(([A-Z0-9]{3})\)(?=;|$)", skills)))
    return out


def read_sot():
    """{'MA.8.NSO.1.3': (statement, [clarification lines])}."""
    out, lines = {}, open(SOT, encoding="utf-8").read().splitlines()
    for i, ln in enumerate(lines):
        m = re.match(r"^\*\*(MA\.\d+\.[A-Z]+\.\d+\.\d+)\*\* — (.+)$", ln)
        if m:
            notes = []
            for nxt in lines[i + 1:]:
                if not nxt.startswith("- "):
                    break
                notes.append(clean_md(nxt[2:]))
            out[m.group(1)] = (clean_md(m.group(2)), notes)
    return out


def read_toc():
    """{'3.6': 'Evaluating Expressions with ...'} from the book's table of contents."""
    out = {}
    for ln in open(TOC, encoding="utf-8"):
        if not re.match(r"^\d+\.\d+ ", ln):
            continue
        head = ln.rstrip("\n").split(" — ")[0]
        waiting = []                               # '4.7 / 4.8 Title': 4.7 shares 4.8's title
        for part in re.split(r" [/·] (?=\d+\.\d+( |$))", head):
            part = (part or "").strip()
            if re.fullmatch(r"\d+\.\d+", part):
                waiting.append(part); continue
            m = re.match(r"^(\d+\.\d+) (.+)$", part)
            if m:
                for n in waiting + [m.group(1)]:
                    out[n] = m.group(2).replace("∅", "").strip()
                waiting = []
    return out


def book_lines(nums, toc):
    """['4.7', '4.8'] -> ['4.7 / 4.8 Title'] when they share a title (as the book prints them)."""
    lines, group = [], []
    for n in [n for n in nums if n in toc]:
        if group and toc[group[-1]] != toc[n]:
            lines.append(" / ".join(group) + " " + toc[group[-1]]); group = []
        group.append(n)
    if group:
        lines.append(" / ".join(group) + " " + toc[group[-1]])
    return lines


def load_specs():
    specs = {}
    for f in sorted(glob.glob(os.path.join(BUILD, "u*", "l*.py"))):
        s = importlib.util.spec_from_file_location("spec", f)
        m = importlib.util.module_from_spec(s)
        with contextlib.redirect_stdout(io.StringIO()):
            s.loader.exec_module(m)
        specs[m.L["code"]] = m.L
    return specs


def load_units():
    units = {}
    for f in sorted(glob.glob(os.path.join(BUILD, "u*", "unit.py"))):
        s = importlib.util.spec_from_file_location("unit", f)
        m = importlib.util.module_from_spec(s)
        with contextlib.redirect_stdout(io.StringIO()):
            s.loader.exec_module(m)
        U = [v for v in vars(m).values() if isinstance(v, dict) and "assessment" in v][0]
        units[U["unit"]] = U
    return units


def stem_of(code, unit):
    """Calendar lesson code -> the file stem its package uses: 3.06+07 -> 3.06, T-A1 -> 3.T1."""
    m = re.fullmatch(r"T-[A-Z](\d)", code)
    if m:
        return f"{unit}.T{m.group(1)}"
    return re.split(r"[+–]", code)[0]


def is_exam(code):
    return bool(re.fullmatch(r"[\d/]+\.X\d", code))


# ---------------------------------------------------------------- links
def gh(path, kind="blob"):
    return f"{REPO}/{kind}/main/" + urllib.parse.quote(os.path.relpath(path, ROOT).replace(os.sep, "/"), safe="/")


def pkg_folder(unit):
    hits = glob.glob(os.path.join(PKG, f"A7 Unit {unit} - *")) if unit else []
    return hits[0] if hits else None


def pfile(unit, sub, name):
    folder = pkg_folder(unit)
    p = os.path.join(folder, sub, name) if folder else None
    return p if p and os.path.exists(p) else None


LINK_COLS = ["Slides (PDF)", "Slides (PowerPoint)", "Teacher Edition", "Lesson Plan", "Question Bank", "Bank key",
             "Bank – Additional", "Additional key", "Independent Set", "Independent key", "Unit folder"]


def row_links(row):
    """{column: (text, url)} for a calendar row; {} when nothing is built for it."""
    U, code = row["unit"], row["code"]
    folder = pkg_folder(U)
    if not folder:
        return {}
    out = {"Unit folder": ("folder", gh(folder, "tree"))}
    if is_exam(code):
        pairs = {"Slides (PDF)": ("Unit Assessment", "PDFs/Assessments", f"A7 {U}  Unit Assessment.pdf"),
                 "Bank key": ("Assessment key", "PDFs/Answer Keys", f"A7 {U}  Unit Assessment Key.pdf")}
    elif code == "spiral":
        pairs = {"Slides (PDF)": ("Unit Review", "PDFs/Assessments", f"A7 {U}  Unit Review.pdf"),
                 "Bank key": ("Review key", "PDFs/Answer Keys", f"A7 {U}  Unit Review Key.pdf")}
    elif code == "flex" or code.startswith("PM"):
        pairs = {}
    else:
        s = stem_of(code, U)
        pairs = {"Slides (PDF)": ("open", "PDFs/Slides", f"A7 {s}  Slides.pdf"),
                 "Slides (PowerPoint)": ("open", "Slides", f"A7 {s}  Slides.pptx"),
                 "Teacher Edition": ("open", "PDFs/Teacher Editions", f"A7 {s}  Teacher Edition.pdf"),
                 "Lesson Plan": ("open", "PDFs/Lesson Plans", f"A7 {s}  Lesson Plan.pdf"),
                 "Question Bank": ("open", "PDFs/Question Banks", f"A7 {s}  Question Bank.pdf"),
                 "Bank key": ("key", "PDFs/Answer Keys", f"A7 {s}  Question Bank Key.pdf"),
                 "Bank – Additional": ("open", "PDFs/Question Banks", f"A7 {s}  Question Bank - Additional.pdf"),
                 "Additional key": ("key", "PDFs/Answer Keys", f"A7 {s}  Question Bank - Additional Key.pdf"),
                 "Independent Set": ("open", "PDFs/Independent Sets", f"A7 {s}  Independent Set.pdf"),
                 "Independent key": ("key", "PDFs/Answer Keys", f"A7 {s}  Independent Set Key.pdf")}
    for col, (text, sub, name) in pairs.items():
        p = pfile(U, sub, name)
        if p:
            out[col] = (text, gh(p))
    if len(out) == 1 and pairs:
        return {}                              # a unit folder exists but this lesson has no documents yet
    return out


# ---------------------------------------------------------------- styles
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
BODY = Font(name=FONT, size=10)
LINK = Font(name=FONT, size=10, color="0563C1", underline="single")
DIM = Font(name=FONT, size=10, color="808080", italic=True)
BOLD = Font(name=FONT, size=10, bold=True)
TITLE = Font(name=FONT, size=14, bold=True, color="1F3864")
SUB = Font(name=FONT, size=11, bold=True, color="1F3864")
EXAM_FILL = PatternFill("solid", fgColor="FCE4D6")
GREY_FILL = PatternFill("solid", fgColor="EDEDED")
NOSCHOOL_FILL = PatternFill("solid", fgColor="D9D9D9")
UNIT_FILLS = [PatternFill("solid", fgColor="FFFFFF"), PatternFill("solid", fgColor="F2F7FC")]
FILL_IN = PatternFill("solid", fgColor="FFF2CC")
BORDER = Border(bottom=Side(style="thin", color="D9D9D9"))
BOX = Border(*(Side(style="thin", color="BFBFBF"),) * 4)


def header(ws, cols, row=1, height=30):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font = HEAD_FONT; cell.fill = HEAD_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = height


def put(ws, r, c, value, font=BODY, link=None, fill=None, wrap=False, fmt=None, align=None, border=BORDER, internal=None):
    cell = ws.cell(row=r, column=c, value=value)
    cell.font = font
    if link:
        cell.hyperlink = link; cell.font = LINK
    if internal:
        cell.hyperlink = Hyperlink(ref=cell.coordinate, location=internal); cell.font = LINK
    if fill:
        cell.fill = fill
    cell.alignment = Alignment(vertical="top", wrap_text=wrap, horizontal=align)
    if border:
        cell.border = border
    if fmt:
        cell.number_format = fmt
    return cell


def note(cell, text, w=360, h=160):
    c = Comment(text, "A7 plan"); c.width = w; c.height = h
    cell.comment = c


def commit():
    try:
        return subprocess.run(["git", "-C", ROOT, "rev-parse", "--short", "HEAD"], capture_output=True, text=True).stdout.strip()
    except Exception:
        return "?"


def bm_codes(cell):
    """'7.NSO.1.1 · 8.NSO.1.3' -> ['MA.7.NSO.1.1', 'MA.8.NSO.1.3']; a range 8.NSO.1.1–1.7 is expanded;
    a cluster (8.AR.2) is returned as it stands, prefixed."""
    out = []
    for b in [x.strip() for x in cell.split("·") if x.strip()]:
        m = re.fullmatch(r"(\d+)\.([A-Z]+)\.(\d+)\.(\d+)–(\d+)\.(\d+)", b)
        if m:
            g, s, a, lo, _, hi = m.groups()
            out += [f"MA.{g}.{s}.{a}.{k}" for k in range(int(lo), int(hi) + 1)]
        else:
            out.append("MA." + b)
    return out


def wording(code, sot):
    if code in sot:
        return f"{code} — {sot[code][0]}"
    if code.startswith("MA.912."):
        return f"{code} — Algebra 1 benchmark (PM3-window bridge); its wording is not in the Grade 8 source file."
    if re.fullmatch(r"MA\.\d+\.[A-Z]+\.\d+", code):
        return f"{code} — every benchmark in this cluster (review day)."
    return f"{code} — (wording not found)"


# ---------------------------------------------------------------- build
def build(out):
    cal = read_calendar(); dues = read_due_sheet(); sot = read_sot(); toc = read_toc()
    specs = load_specs(); units = load_units()
    wb = Workbook()
    ws_today = wb.active; ws_today.title = "Today"
    ws = wb.create_sheet("Scope & Sequence")

    # ------------------------------------------------------------ Scope & Sequence
    cols = (["Date", "Day", "Unit", "Lesson", "What is taught"] + LINK_COLS +
            ["IXL skills — all required, code", "IXL due", "Learning target", "Essential question", "Vocabulary",
             "Benchmark", "Benchmark wording", "MTRs (tap for where they show up)", "Closure question (board 9)",
             "Closure answer", "Math Nation lesson(s)"])
    C = {name: i for i, name in enumerate(cols, 1)}           # column numbers by header
    L_ = {name: get_column_letter(i) for name, i in C.items()}  # column letters by header
    header(ws, cols)
    first, last = 2, 1 + len(cal)
    rng = lambda name: f"'Scope & Sequence'!${L_[name]}${first}:${L_[name]}${last}"
    scope_row = {}
    links_by_row = []
    for r, row in enumerate(cal, first):
        scope_row[row["date"]] = r
        code, U = row["code"], row["unit"]
        links = row_links(row); links_by_row.append(links)
        exam = is_exam(code)
        fill = EXAM_FILL if exam else GREY_FILL if code in ("spiral", "flex") or code.startswith("PM") else UNIT_FILLS[(U or 0) % 2]
        L = specs.get(stem_of(code, U)) if not exam and code not in ("spiral", "flex") and not code.startswith("PM") else None
        put(ws, r, C["Date"], row["date"], fill=fill, fmt="ddd d mmm yyyy")
        put(ws, r, C["Day"], row["day"], fill=fill)
        put(ws, r, C["Unit"], U, fill=fill, align="center")
        put(ws, r, C["Lesson"], code, fill=fill, font=BOLD)
        put(ws, r, C["What is taught"], row["taught"], fill=fill, wrap=True)
        for col in LINK_COLS:
            hit = links.get(col)
            if hit:
                put(ws, r, C[col], hit[0], link=hit[1], fill=fill)
            elif col == "Slides (PDF)" and U and not pkg_folder(U) and not code.startswith("PM"):
                put(ws, r, C[col], "not built yet", font=DIM, fill=fill)
            else:
                put(ws, r, C[col], None, fill=fill)
        put(ws, r, C["IXL skills — all required, code"], row["ixl"] or None, fill=fill, wrap=True)
        put(ws, r, C["IXL due"], row["due"], fill=fill, fmt="ddd d mmm")
        codes = bm_codes(row["bm"]) if row["bm"] else []
        put(ws, r, C["Benchmark"], row["bm"] or None, fill=fill, wrap=True)
        put(ws, r, C["Benchmark wording"], "\n\n".join(wording(c, sot) for c in codes) or None, fill=fill, wrap=True)
        mn = book_lines(SC.lessons_for(code), toc) if not exam else []
        put(ws, r, C["Math Nation lesson(s)"], "\n".join(mn) or None, fill=fill, wrap=True)
        if L:
            put(ws, r, C["Learning target"], tex2text.text(L["target"]), fill=fill, wrap=True)
            put(ws, r, C["Essential question"], tex2text.text(L["essential"]), fill=fill, wrap=True)
            put(ws, r, C["Vocabulary"], ", ".join(t for t, _ in L["vocab"]), fill=fill, wrap=True)
            mtr_cell = put(ws, r, C["MTRs (tap for where they show up)"],
                           "\n".join(f"MA.K12.{k} {MTR[k]}" for k, _ in L["mtr"]), fill=fill, wrap=True)
            note(mtr_cell, "\n\n".join(f"{k}: {tex2text.text(v)}" for k, v in L["mtr"]), 420, 220)
            b9 = L["whiteboard"][8]
            q = put(ws, r, C["Closure question (board 9)"], tex2text.text(b9["qtext"]), fill=fill, wrap=True)
            note(q, "What to look for: " + tex2text.text(L["closure"]), 420, 160)
            put(ws, r, C["Closure answer"], tex2text.text(b9["answer"]), fill=fill, wrap=True)
        else:
            for col in ("Learning target", "Essential question", "Vocabulary", "MTRs (tap for where they show up)",
                        "Closure question (board 9)", "Closure answer"):
                put(ws, r, C[col], None, fill=fill)
    widths = {"Date": 15, "Day": 12, "Unit": 5, "Lesson": 10, "What is taught": 40, "IXL skills — all required, code": 44,
              "IXL due": 11, "Learning target": 44, "Essential question": 36, "Vocabulary": 30, "Benchmark": 16,
              "Benchmark wording": 60, "MTRs (tap for where they show up)": 46, "Closure question (board 9)": 50,
              "Closure answer": 30, "Math Nation lesson(s)": 42}
    for name, i in C.items():
        ws.column_dimensions[get_column_letter(i)].width = widths.get(name, 11 if name != "Unit folder" else 9)
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last}"

    # ------------------------------------------------------------ Links (hidden; feeds Today's buttons)
    wl = wb.create_sheet("Links")
    wl.cell(row=1, column=1, value="Date")
    for j, col in enumerate(LINK_COLS, 2):
        wl.cell(row=1, column=j, value=col)
    for r, (row, links) in enumerate(zip(cal, links_by_row), first):
        wl.cell(row=r, column=1, value=row["date"]).number_format = "d mmm yyyy"
        for j, col in enumerate(LINK_COLS, 2):
            if col in links:
                wl.cell(row=r, column=j, value=links[col][1])
    for row in wl.iter_rows():
        for c in row:
            c.font = BODY
    wl.sheet_state = "hidden"
    link_rng = lambda col: f"Links!${get_column_letter(LINK_COLS.index(col) + 2)}${first}:${get_column_letter(LINK_COLS.index(col) + 2)}${last}"

    # ------------------------------------------------------------ IXL tracker
    wi = wb.create_sheet("IXL tracker")
    wi.cell(row=1, column=1, value="IXL assignments, 2026–27").font = TITLE
    wi.cell(row=2, column=1, value="Every skill listed is required, to a SmartScore of 67. Fill in the two yellow columns: "
                                    "type ✓ in 'Set up in IXL' once the assignment exists in IXL, and ✓ in 'Scores checked' "
                                    "after you have looked at the scores (for example, row 1 done: ✓ and ✓, Status reads 'done'). "
                                    "Status works itself out from today's date.").font = DIM
    wi.merge_cells(start_row=2, start_column=1, end_row=2, end_column=11)
    wi.cell(row=2, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    wi.row_dimensions[2].height = 42
    icols = ["#", "Assigned", "Lesson(s)", "Skills — all required", "Codes to type into IXL", "SmartScore", "Due",
             "Set up in IXL", "Scores checked", "Status", "Notes"]
    header(wi, icols, row=4)
    for k, a in enumerate(dues, 1):
        r = 4 + k
        put(wi, r, 1, k, align="center")
        put(wi, r, 2, a["assigned"], fmt="ddd d mmm")
        lr = scope_row.get(a["assigned"])
        put(wi, r, 3, a["lessons"], wrap=True, internal=f"'Scope & Sequence'!D{lr}" if lr else None)
        put(wi, r, 4, a["skills"], wrap=True)
        put(wi, r, 5, ", ".join(a["codes"]))
        put(wi, r, 6, 67, align="center")
        put(wi, r, 7, a["due"], fmt="ddd d mmm")
        put(wi, r, 8, None, fill=FILL_IN, align="center")
        put(wi, r, 9, None, fill=FILL_IN, align="center")
        put(wi, r, 10, f'=IF(I{r}="✓","done",IF(TODAY()<B{r},"upcoming",IF(TODAY()<G{r},"open","due — check scores")))')
        put(wi, r, 11, None, fill=FILL_IN, wrap=True)
    ilast = 4 + len(dues)
    dv = DataValidation(type="list", formula1='"✓"', allow_blank=True, showErrorMessage=False)
    wi.add_data_validation(dv); dv.add(f"H5:I{ilast}")
    wi.conditional_formatting.add(f"J5:J{ilast}", FormulaRule(formula=[f'$J5="done"'], fill=PatternFill("solid", fgColor="E2EFDA")))
    wi.conditional_formatting.add(f"J5:J{ilast}", FormulaRule(formula=[f'$J5="due — check scores"'], fill=PatternFill("solid", fgColor="F8CBAD")))
    wi.conditional_formatting.add(f"J5:J{ilast}", FormulaRule(formula=[f'$J5="open"'], fill=PatternFill("solid", fgColor="FFF2CC")))
    for i, w in enumerate([5, 13, 40, 60, 18, 10, 12, 12, 12, 18, 30], 1):
        wi.column_dimensions[get_column_letter(i)].width = w
    wi.freeze_panes = "C5"
    wi.auto_filter.ref = f"A4:K{ilast}"

    # ------------------------------------------------------------ Unit tests
    wt = wb.create_sheet("Unit tests")
    wt.cell(row=1, column=1, value="Unit tests and progress monitoring").font = TITLE
    tcols = ["Unit", "Exam day 1", "Exam day 2", "Points", "Paper", "Key", "Unit Review (practice)", "Review key", "Reference Sheet"]
    header(wt, tcols, row=3)
    exams = {}
    for row in cal:
        if is_exam(row["code"]):
            exams.setdefault(row["code"].split(".X")[0], []).append(row["date"])
    r = 4
    total_cell = {}
    test_rows = [("1", None), ("2", None)] + [(k, v) for k, v in exams.items()]
    points_row = {}
    for label, dates in test_rows:
        u = int(label.split("/")[-1])
        put(wt, r, 1, f"Unit {label}", font=BOLD)
        if dates:
            put(wt, r, 2, dates[0], fmt="ddd d mmm"); put(wt, r, 3, dates[1], fmt="ddd d mmm")
        else:
            put(wt, r, 2, "before 23 Sep", font=DIM); put(wt, r, 3, None)
        points_row[u] = r
        put(wt, r, 4, None, align="center")
        for j, (sub, name) in enumerate([("PDFs/Assessments", f"A7 {u}  Unit Assessment.pdf"),
                                          ("PDFs/Answer Keys", f"A7 {u}  Unit Assessment Key.pdf"),
                                          ("PDFs/Assessments", f"A7 {u}  Unit Review.pdf"),
                                          ("PDFs/Answer Keys", f"A7 {u}  Unit Review Key.pdf"),
                                          ("PDFs/Handouts", f"A7 {u}  Reference Sheet.pdf")], 5):
            p = pfile(u, sub, name)
            if p:
                put(wt, r, j, "open", link=gh(p))
            else:
                put(wt, r, j, "not built yet" if (j == 5 and not pkg_folder(u)) else None, font=DIM)
        r += 1
    pm = [x for x in cal if x["code"].startswith("PM2")]
    put(wt, r, 1, "PM2", font=BOLD)
    put(wt, r, 2, pm[0]["date"], fmt="ddd d mmm"); put(wt, r, 3, pm[-1]["date"], fmt="ddd d mmm")
    put(wt, r, 4, None); put(wt, r, 5, "review days; the PM2 test date is set by the school", font=DIM)
    wt.merge_cells(start_row=r, start_column=5, end_row=r, end_column=9); r += 1
    put(wt, r, 1, "PM3", font=BOLD)
    put(wt, r, 2, datetime.date(2027, 5, 3), fmt="ddd d mmm"); put(wt, r, 3, SC.end, fmt="ddd d mmm")
    put(wt, r, 4, None); put(wt, r, 5, "the PM3 window; test date TBD by the school — the whole window is free of new content", font=DIM)
    wt.merge_cells(start_row=r, start_column=5, end_row=r, end_column=9); r += 2
    # item maps
    item_cols = ["Question", "Parts", "Points", "Section", "Benchmark", "Format", "Transfer item"]
    for u in sorted(units):
        U = units[u]
        total, per_bm, per_sec, transfer = unitbuild.assessment_ledger(U)
        put(wt, r, 1, f"Unit {u} test — question map", font=SUB, border=None); r += 1
        put(wt, r, 1, "One paper over two periods; one point per lettered part. Transfer items are not on the practice test.", font=DIM, border=None)
        r += 1
        header(wt, item_cols, row=r, height=18)
        first_item = r + 1
        n = 0
        for sec in U["assessment"]["sections"]:
            for it in sec["items"]:
                if it.get("heading"):
                    continue
                n += 1; r += 1
                parts = it.get("parts") or []
                pts = unitbuild._points(it); pts = sum(pts) if isinstance(pts, list) else pts
                put(wt, r, 1, n, align="center")
                put(wt, r, 2, f"{parts[0]['label']}–{parts[-1]['label']}" if len(parts) > 1 else (parts[0]["label"] if parts else None), align="center")
                put(wt, r, 3, pts, align="center")
                put(wt, r, 4, sec["title"])
                put(wt, r, 5, sec["benchmark"])
                put(wt, r, 6, "Multiple choice" if it.get("choices") else ("Parts" if parts else "Written"))
                put(wt, r, 7, "yes" if it.get("transfer") else None, align="center")
        last_item = r
        r += 1
        put(wt, r, 2, "Total", font=BOLD); tc = put(wt, r, 3, f"=SUM(C{first_item}:C{last_item})", font=BOLD, align="center")
        total_cell[u] = f"C{r}"
        # per-benchmark summary to the right (columns I-K), clear of the Standards tab's SUMIF columns
        sr = first_item - 1
        for j, h in enumerate(["Benchmark", "Points", "Share of the test"], 9):
            c = wt.cell(row=sr, column=j, value=h); c.font = HEAD_FONT; c.fill = HEAD_FILL
        for k, bm in enumerate(U["assessment"]["tracker_order"], 1):
            put(wt, sr + k, 9, bm)
            put(wt, sr + k, 10, f"=SUMIF($E${first_item}:$E${last_item},I{sr + k},$C${first_item}:$C${last_item})", align="center")
            put(wt, sr + k, 11, f"=J{sr + k}/$C${r}", fmt="0%", align="center")
        wt.cell(row=points_row[u], column=4, value=f"={total_cell[u]}")
        r += 3
    for i, w in enumerate([12, 12, 12, 9, 34, 14, 22, 12, 16, 9, 15], 1):
        wt.column_dimensions[get_column_letter(i)].width = w
    wt.freeze_panes = "A4"

    # ------------------------------------------------------------ Standards
    wsd = wb.create_sheet("Standards")
    scols = ["Benchmark", "Grade", "Wording", "Clarifications and limits", "Unit(s)", "Taught on", "First day", "Last day",
             "Reviewed on", "Unit test points", "IXL skill codes"]
    header(wsd, scols)
    taught, reviewed, ixl_by, units_by = {}, {}, {}, {}
    for row in cal:
        if not row["bm"]:
            continue
        codes = bm_codes(row["bm"])
        review = row["code"].startswith("PM")
        for c in sot.keys() | {x for x in codes if x.startswith("MA.912")}:
            hit = c in codes or any(re.fullmatch(r"MA\.\d+\.[A-Z]+\.\d+", x) and c.startswith(x + ".") for x in codes)
            if not hit:
                continue
            if review:
                reviewed.setdefault(c, []).append(row)
            else:
                taught.setdefault(c, []).append(row)
                if row["unit"]:
                    units_by.setdefault(c, set()).add(row["unit"])
                ixl_by.setdefault(c, [])
                for code in re.findall(r"— ([A-Z0-9]{3})(?=;|$)", row["ixl"]):
                    if code not in ixl_by[c]:
                        ixl_by[c].append(code)
    early = {f"MA.{b}": u for u, bs in EARLY_UNITS.items() for b in bs}
    allcodes = sorted(set(sot) | set(taught) | set(reviewed),
                      key=lambda c: (0 if c in early else 1, early.get(c, 0),
                                     min(x["date"] for x in taught[c]) if c in taught else datetime.date(2099, 1, 1), c))
    for r, c in enumerate(allcodes, 2):
        stmt, notes = sot.get(c, (None, []))
        g = c.split(".")[1]
        put(wsd, r, 1, c, font=BOLD)
        put(wsd, r, 2, "Algebra 1" if g == "912" else int(g), align="center")
        put(wsd, r, 3, stmt or wording(c, sot).split(" — ", 1)[1], wrap=True)
        put(wsd, r, 4, "\n".join(notes) or None, wrap=True)
        us = sorted(units_by.get(c, set()) | ({early[c]} if c in early else set()))
        put(wsd, r, 5, ", ".join(str(u) for u in us) or ("PM3 window" if g == "912" else None), align="center")
        days = taught.get(c, [])
        txt = "; ".join(f"{x['code']} ({x['date'].strftime('%a')} {x['date'].day} {x['date'].strftime('%b')})" for x in days)
        if c in early:
            txt = (f"Unit {early[c]} (before 23 Sep)" + ("; " + txt if txt else ""))
        put(wsd, r, 6, txt or None, wrap=True)
        put(wsd, r, 7, min(x["date"] for x in days) if days else None, fmt="d mmm yyyy")
        put(wsd, r, 8, max(x["date"] for x in days) if days else None, fmt="d mmm yyyy")
        rv = reviewed.get(c, [])
        put(wsd, r, 9, "; ".join(f"{x['code']} ({x['date'].day} {x['date'].strftime('%b')})" for x in rv) or None, wrap=True)
        put(wsd, r, 10, f"=IF(SUMIF('Unit tests'!$E:$E,A{r},'Unit tests'!$C:$C)=0,\"\",SUMIF('Unit tests'!$E:$E,A{r},'Unit tests'!$C:$C))", align="center")
        put(wsd, r, 11, ", ".join(ixl_by.get(c, [])) or None, wrap=True)
    for i, w in enumerate([15, 9, 60, 50, 8, 40, 12, 12, 22, 10, 22], 1):
        wsd.column_dimensions[get_column_letter(i)].width = w
    wsd.freeze_panes = "B2"
    wsd.auto_filter.ref = f"A1:K{1 + len(allcodes)}"

    # ------------------------------------------------------------ Month calendar
    wm = wb.create_sheet("Month calendar")
    by_date = {x["date"]: x for x in cal}
    days = set(SC.days)
    r = 1
    month = datetime.date(2026, 9, 1)
    wd_names = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]
    breaks = []
    while month <= datetime.date(2027, 5, 1):
        nxt = (month.replace(day=28) + datetime.timedelta(days=4)).replace(day=1)
        c = wm.cell(row=r, column=1, value=month.strftime("%B %Y")); c.font = TITLE
        wm.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5); r += 1
        for j, n in enumerate(wd_names, 1):
            c = wm.cell(row=r, column=j, value=n); c.font = HEAD_FONT; c.fill = HEAD_FILL; c.alignment = Alignment(horizontal="center")
        r += 1
        d = month - datetime.timedelta(days=month.weekday())
        if month.month == 9:
            d = datetime.date(2026, 9, 21)                  # the plan starts in the week of 21 Sep
        while d < nxt:
            if d.weekday() == 0:
                for j in range(5):
                    day = d + datetime.timedelta(days=j)
                    cell = wm.cell(row=r, column=j + 1)
                    cell.border = BOX; cell.alignment = Alignment(wrap_text=True, vertical="top")
                    cell.font = BODY
                    if day.month != month.month or day > SC.end:
                        continue                    # outside the month, or after the last day of school
                    if day in by_date:
                        x = by_date[day]
                        label = x["taught"] if not x["code"].startswith("PM3") else "PM3 window"
                        cell.value = f"{day.day}   {x['code']}\n{label}"
                        cell.fill = EXAM_FILL if is_exam(x["code"]) else GREY_FILL if x["code"] in ("spiral", "flex") or x["code"].startswith("PM") else PatternFill(fill_type=None)
                        if x["code"] not in ("flex",) and not x["code"].startswith("PM") and not is_exam(x["code"]) and x["code"] != "spiral":
                            cell.fill = UNIT_FILLS[(x["unit"] or 0) % 2]
                    elif day < START:
                        cell.value = f"{day.day}\nUnits 1–2"; cell.fill = GREY_FILL; cell.font = DIM
                    elif day in days and day >= datetime.date(2027, 5, 3):
                        cell.value = f"{day.day}\nPM3 window"; cell.fill = GREY_FILL
                    elif day in SC.hol:
                        cell.value = f"{day.day}\nNo school"; cell.fill = NOSCHOOL_FILL; cell.font = DIM
                    else:
                        cell.value = f"{day.day}\n(not in the plan)"; cell.font = DIM
                wm.row_dimensions[r].height = 66
                r += 1
            d += datetime.timedelta(days=7 - d.weekday()) if d.weekday() else datetime.timedelta(days=7)
        breaks.append(r); r += 1
        month = nxt
    for b in breaks[:-1]:
        wm.row_breaks.append(Break(id=b))
    for j in range(1, 6):
        wm.column_dimensions[get_column_letter(j)].width = 28
    wm.page_setup.orientation = "landscape"
    wm.page_setup.fitToWidth = 1; wm.page_setup.fitToHeight = 0
    wm.sheet_properties.pageSetUpPr.fitToPage = True
    wm.print_options.horizontalCentered = True

    # ------------------------------------------------------------ Vocabulary
    wv = wb.create_sheet("Vocabulary")
    header(wv, ["Unit", "Lesson", "First day", "Term", "Definition"])
    r = 2
    firsts = {}
    for row in cal:
        s = stem_of(row["code"], row["unit"]) if row["unit"] else None
        if s in specs and s not in firsts:
            firsts[s] = row
    for s, row in sorted(firsts.items(), key=lambda kv: kv[1]["date"]):
        for term, d in specs[s]["vocab"]:
            put(wv, r, 1, row["unit"], align="center"); put(wv, r, 2, row["code"], font=BOLD)
            put(wv, r, 3, row["date"], fmt="ddd d mmm", internal=f"'Scope & Sequence'!A{scope_row[row['date']]}")
            put(wv, r, 4, term, font=BOLD, wrap=True); put(wv, r, 5, tex2text.text(d), wrap=True)
            r += 1
    for i, w in enumerate([6, 10, 13, 28, 90], 1):
        wv.column_dimensions[get_column_letter(i)].width = w
    wv.freeze_panes = "D2"
    wv.auto_filter.ref = f"A1:E{r - 1}"

    # ------------------------------------------------------------ Units
    wu = wb.create_sheet("Units")
    ucols = ["Unit", "Title", "First day", "Last day", "Periods", "Status", "Package folder", "START HERE", "Reference Sheet",
             "Unit Review", "Review key", "Unit Assessment", "Assessment key", "Slides folder"]
    header(wu, ucols)
    r = 2
    for u in range(1, 18):
        folder = pkg_folder(u)
        ds = [x["date"] for x in cal if x["unit"] == u]
        put(wu, r, 1, u, align="center", font=BOLD)
        put(wu, r, 2, UNIT_TITLES[u], wrap=True)
        if ds:
            put(wu, r, 3, min(ds), fmt="d mmm yyyy"); put(wu, r, 4, max(ds), fmt="d mmm yyyy")
            put(wu, r, 5, f"=COUNTIF({rng('Unit')},A{r})", align="center")
        else:
            put(wu, r, 3, "before 23 Sep", font=DIM); put(wu, r, 4, None); put(wu, r, 5, None)
        if folder:
            put(wu, r, 6, "built, audited 21 Sep 2026" if u in (3, 4) else "built (before the spec system; not covered by the checks)", wrap=True)
            put(wu, r, 7, "folder", link=gh(folder, "tree"))
            sh = os.path.join(folder, "00 - START HERE.md")
            put(wu, r, 8, "open", link=gh(sh)) if os.path.exists(sh) else put(wu, r, 8, None)
            for j, (sub, name) in enumerate([("PDFs/Handouts", f"A7 {u}  Reference Sheet.pdf"), ("PDFs/Assessments", f"A7 {u}  Unit Review.pdf"),
                                              ("PDFs/Answer Keys", f"A7 {u}  Unit Review Key.pdf"), ("PDFs/Assessments", f"A7 {u}  Unit Assessment.pdf"),
                                              ("PDFs/Answer Keys", f"A7 {u}  Unit Assessment Key.pdf")], 9):
                p = pfile(u, sub, name)
                put(wu, r, j, "open", link=gh(p)) if p else put(wu, r, j, None)
            put(wu, r, 14, "folder", link=gh(os.path.join(folder, "PDFs", "Slides"), "tree"))
        else:
            put(wu, r, 6, "not built yet", font=DIM)
            for j in range(7, 15):
                put(wu, r, j, None)
        r += 1
    put(wu, r + 1, 1, "Periods counts this unit's rows on the Scope & Sequence tab (exam, spiral and flex days included).", font=DIM, border=None)
    for i, w in enumerate([6, 40, 13, 13, 9, 30, 10, 11, 11, 10, 10, 11, 12, 10], 1):
        wu.column_dimensions[get_column_letter(i)].width = w
    wu.freeze_panes = "C2"

    # ------------------------------------------------------------ Units 1–2 decks
    wd = wb.create_sheet("Units 1–2 decks")
    header(wd, ["Unit", "Lesson", "Slides (PDF)", "Teacher Edition", "Question Bank", "Question Bank – Additional", "Answer keys (PDF)"])
    r = 2
    for u in (1, 2):
        folder = pkg_folder(u)
        if not folder:
            continue
        for pdf in sorted(glob.glob(os.path.join(folder, "PDFs", "Slides", "A7 *  Slides.pdf"))):
            s = re.match(r"A7 (\S+)  Slides\.pdf", os.path.basename(pdf)).group(1)
            put(wd, r, 1, u, align="center", font=BOLD); put(wd, r, 2, s, font=BOLD)
            put(wd, r, 3, "open", link=gh(pdf))
            for j, (sub, name) in enumerate([("PDFs/Teacher Editions", f"A7 {s}  Teacher Edition.pdf"),
                                              ("PDFs/Question Banks", f"A7 {s}  Question Bank.pdf"),
                                              ("PDFs/Question Banks", f"A7 {s}  Question Bank - Additional.pdf")], 4):
                p = pfile(u, sub, name)
                put(wd, r, j, "open", link=gh(p)) if p else put(wd, r, j, None)
            put(wd, r, 7, "folder", link=gh(os.path.join(folder, "PDFs", "Answer Keys"), "tree"))
            r += 1
    for i, w in enumerate([6, 9, 13, 16, 14, 24, 16], 1):
        wd.column_dimensions[get_column_letter(i)].width = w
    wd.freeze_panes = "C2"

    # ------------------------------------------------------------ Today
    t = ws_today
    t.column_dimensions["A"].width = 24; t.column_dimensions["B"].width = 70; t.column_dimensions["C"].width = 16
    t.cell(row=1, column=1, value="A7 — today").font = TITLE
    c = t.cell(row=1, column=2, value="=TODAY()"); c.number_format = "dddd d mmmm yyyy"; c.font = SUB
    N = len(cal)
    # helper cells, at the bottom: which plan row to show, and the last day of school. The last row
    # is the PM3 window (3-28 May): on any day in it, Today shows that row.
    HELP = 42
    put(t, HELP, 1, "Plan row shown", font=DIM, border=None)
    put(t, HELP, 2, f"=IF(TODAY()>$B${HELP + 1},{N}+1,MIN(COUNTIF({rng('Date')},\"<\"&TODAY())+1,{N}))", font=DIM, border=None, align="left")
    put(t, HELP + 1, 1, "Last day of school", font=DIM, border=None)
    put(t, HELP + 1, 2, SC.end, font=DIM, border=None, fmt="d mmm yyyy", align="left")
    idx = f"$B${HELP}"
    def at(name):
        return f"INDEX({rng(name)},{idx})"
    put(t, 3, 1, "Showing", font=BOLD, border=None)
    put(t, 3, 2, f'=IF({idx}>{N},"The school year has ended.",IF({at("Date")}=TODAY(),"Today\'s class",IF({at("Date")}<TODAY(),"The PM3 window (3–28 May)","Next class: "&TEXT({at("Date")},"dddd d mmmm"))))', font=SUB, border=None)
    fields = [("Lesson", f'=IF({idx}>{N},"",{at("Lesson")}&"   "&{at("What is taught")})'),
              ("Learning target", f'=IF({idx}>{N},"",{at("Learning target")}&"")'),
              ("Essential question", f'=IF({idx}>{N},"",{at("Essential question")}&"")'),
              ("Vocabulary", f'=IF({idx}>{N},"",{at("Vocabulary")}&"")'),
              ("Benchmark", f'=IF({idx}>{N},"",{at("Benchmark")}&"")'),
              ("IXL assigned", f'=IF({idx}>{N},"",IF({at("IXL skills — all required, code")}&""="","none",{at("IXL skills — all required, code")}&"   — due "&TEXT({at("IXL due")},"ddd d mmm")))'),
              ("IXL due at the start of class", f"=IF({idx}>{N},\"\",IFERROR(INDEX('IXL tracker'!$D$5:$D${ilast},MATCH({at('Date')},'IXL tracker'!$G$5:$G${ilast},0)),\"nothing due\"))"),
              ("Closure question", f'=IF({idx}>{N},"",{at("Closure question (board 9)")}&"")')]
    r = 5
    for lab, f in fields:
        put(t, r, 1, lab, font=BOLD); put(t, r, 2, f, wrap=True); r += 1
    r += 1
    put(t, r, 1, "Open", font=SUB, border=None); r += 1
    for col in LINK_COLS:
        put(t, r, 1, col)
        put(t, r, 2, f'=IF({idx}>{N},"—",IF(INDEX({link_rng(col)},{idx})&""="","—",HYPERLINK(INDEX({link_rng(col)},{idx}),{at(col)}&"")))', font=LINK)
        r += 1
    r += 1
    put(t, r, 1, "Coming up", font=SUB, border=None); r += 1
    header(t, ["Date", "Lesson", "Slides"], row=r, height=18); r += 1
    for k in range(1, 6):
        i = f"({idx}+{k})"
        put(t, r, 1, f"=IF({i}>{N},\"\",INDEX({rng('Date')},{i}))", fmt="ddd d mmm", align="left")
        put(t, r, 2, f'=IF({i}>{N},"",INDEX({rng("Lesson")},{i})&"   "&INDEX({rng("What is taught")},{i}))', wrap=True)
        put(t, r, 3, f'=IF({i}>{N},"",IF(INDEX({link_rng("Slides (PDF)")},{i})&""="","—",HYPERLINK(INDEX({link_rng("Slides (PDF)")},{i}),INDEX({rng("Slides (PDF)")},{i})&"")))', font=LINK)
        r += 1
    r += 1
    put(t, r, 1, "This tab follows today's date: on a class day it shows that day; on a weekend or holiday, the next class day. "
                 "The links open the file on GitHub (sign in as croix18 if it asks).", font=DIM, border=None)
    t.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    t.cell(row=r, column=1).alignment = Alignment(wrap_text=True, vertical="top"); t.row_dimensions[r].height = 30
    t.page_setup.orientation = "portrait"; t.page_setup.fitToWidth = 1; t.page_setup.fitToHeight = 0
    t.sheet_properties.pageSetUpPr.fitToPage = True

    # ------------------------------------------------------------ About
    wa = wb.create_sheet("About")
    lines = [
        ("What this is", BOLD),
        ("The 2026–27 plan for Grade 7 Accelerated (course 1205050) with everything for each day in one place: one row per class day from 23 September to 30 April, then the PM3 window (3–28 May).", BODY),
        ("", BODY),
        ("The tabs", BOLD),
        ("Today — the next class day (today's, on a class day), with its target, vocabulary, IXL and a tap-to-open link for every document; then the five days after it.", BODY),
        ("Scope & Sequence — every class day: the documents (slides, Teacher Edition, plan, question banks, independent set, and each one's answer key), IXL, learning target, essential question, vocabulary, benchmark wording, MTRs, the closure question and its answer, and the Math Nation lessons the day covers. Tap an MTR or closure cell for the note behind it.", BODY),
        ("On an exam day the Slides column opens the test paper and the Bank key column its key; on a spiral day they open the Unit Review (which goes home as practice) and its key.", BODY),
        ("IXL tracker — all 82 assignments with codes and dates. The two yellow columns are yours to fill in; Status follows today's date.", BODY),
        ("Standards — every benchmark in the year: its wording and limits, the days that teach it, the review days, its points on each unit test, and its IXL skills.", BODY),
        ("Unit tests — every test date, the PM windows, and a question-by-question map (benchmark and points) for each built test.", BODY),
        ("Month calendar — a printable grid, one month per page.", BODY),
        ("Vocabulary — every term and definition, by lesson.", BODY),
        ("Units, Units 1–2 decks — unit-level documents, and the lesson decks for Units 1 and 2.", BODY),
        ("", BODY),
        ("Good to know", BOLD),
        ("Links open the file on GitHub. A PDF opens right in the browser; a PowerPoint file shows a download button instead. If GitHub asks you to sign in, use your croix18 account.", BODY),
        ("'not built yet' means the unit's documents do not exist yet. Units 5 onward fill in as they are built; this file is regenerated from the plan each time.", BODY),
        ("IXL: every skill listed is required, to a SmartScore of 67, due at the start of the next class (ruling 28). Lessons in a row that use the same skills are one assignment, due after the last of them — the IXL due column already shows that date.", BODY),
        ("'No school' days are the holidays in the plan (tools/scope_calendar.py), taken from the board-approved calendar — check them against the school's copy.", BODY),
        ("", BODY),
        ("Where it comes from", BOLD),
        (f"Generated by tools/master_sheet.py at commit {commit()} on {datetime.date.today().strftime('%d %B %Y')}, from the calendar table, the IXL due-date sheet, the Source of Truth (benchmark wording), Florida's B.E.S.T. standards (MTR titles), the Math Nation table of contents, the lesson and unit specs, and the package folders. The plan is edited in scope_calendar.py; this file is never edited by hand — except the yellow cells in the IXL tracker, which are yours.", BODY),
    ]
    for i, (text, f) in enumerate(lines, 1):
        c = wa.cell(row=i, column=1, value=text); c.font = f; c.alignment = Alignment(wrap_text=True, vertical="top")
    wa.column_dimensions["A"].width = 120

    # order: Today first, Links last (hidden)
    order = ["Today", "Scope & Sequence", "IXL tracker", "Standards", "Unit tests", "Month calendar", "Vocabulary",
             "Units", "Units 1–2 decks", "About", "Links"]
    wb._sheets = [wb[n] for n in order]
    wb.active = 0
    wb.save(out)
    return len(cal)


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(REF, "A7 Master Sheet 2026-27.xlsx")
    n = build(out)
    print(f"wrote {out}: {n} class-day rows")
