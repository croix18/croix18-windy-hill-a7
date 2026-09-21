#!/usr/bin/env python3
"""Check `A7 Master Sheet 2026-27.xlsx` against its sources, independently of the generator.

    python3 tools/master_sheet.py && python3 <xlsx skill>/scripts/recalc.py "<xlsx>" && python3 tools/check_master_sheet.py

Parses every source its own way (not with master_sheet.py's helpers) and reads the recalculated
values: every row, date, weekday and due date against the calendar; every link against the git
tree (and each deck's title slide against its row); every built lesson's target, question, vocabulary,
closure and MTRs against its spec; the benchmark wording against the Source of Truth; the Math
Nation lessons against the book's contents; the IXL tracker against the due-date sheet; each unit
test's question map and points against its spec; the month grid against the school days; no
LaTeX left anywhere. With Florida's B.E.S.T. standards text present (outside the repo, at
/root/best), also checks the MTR titles and that every Source of Truth statement is complete and
verbatim. Exit status 1 on any problem."""
import os, re, sys, glob, subprocess, datetime, urllib.parse, importlib.util, collections, io, contextlib
import pymupdf
from openpyxl import load_workbook
ROOT = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
XLSX = sys.argv[1] if len(sys.argv) > 1 else f"{ROOT}/a7/reference/A7 Master Sheet 2026-27.xlsx"
REF = f"{ROOT}/a7/reference"
BASE = "https://github.com/croix18/croix18-windy-hill-a7/"
P = []
def bad(*a): P.append(" ".join(str(x) for x in a))
tracked = set(subprocess.run(["git", "-C", ROOT, "ls-tree", "-r", "HEAD", "--name-only"], capture_output=True, text=True).stdout.splitlines())
def target(url):
    if not url.startswith(BASE): bad("foreign link", url); return None, None
    kind, rest = url[len(BASE):].split("/main/", 1); path = urllib.parse.unquote(rest)
    if kind == "blob" and path not in tracked: bad("blob not tracked:", path)
    if kind == "tree" and not any(t.startswith(path + "/") for t in tracked): bad("tree empty:", path)
    return kind, path
MON = {m: i for i, m in enumerate("Jan Feb Mar Apr May Jun Jul Aug Sep Oct Nov Dec".split(), 1)}
def d(s):
    m, day = s.split()[-2:]; return datetime.date(2026 if MON[m] >= 7 else 2027, MON[m], int(day))
def D(v): return v.date() if isinstance(v, datetime.datetime) else v
def norm(s): return re.sub(r"[^a-z0-9]", "", s.lower())

# ---------------- sources, parsed independently
cal = [[c.strip() for c in l.strip().strip("|").split("|")] for l in open(f"{REF}/A7 SCOPE AND SEQUENCE 2026-27.md", encoding="utf-8") if l.startswith("| ") and not l.startswith("| Date")]
due = [[c.strip() for c in l.strip().strip("|").split("|")] for l in open(f"{REF}/A7 IXL DUE DATES 2026-27.md", encoding="utf-8") if l.startswith("| ") and not l.startswith("| Assigned")]
sot_txt = open(f"{REF}/Florida BEST Grade 8 - Source of Truth.md", encoding="utf-8").read()
SOT = {}
for m in re.finditer(r"^\*\*(MA\.[\w.]+)\*\* — (.+)$", sot_txt, re.M):
    SOT[m.group(1)] = re.sub(r"\*", "", m.group(2)).strip()
toc_txt = open(f"{REF}/MATH NATION A7 BOOK - table of contents and IXL plan.md", encoding="utf-8").read()
STD = "/root/best/std/mathbeststandardsfinal/mathbeststandardsfinal.txt"
std = open(STD, encoding="utf-8", errors="replace").read() if os.path.exists(STD) else None
std_n = norm(std.replace(" - ", "-")) if std else None
sys.path.insert(0, os.path.join(ROOT, "tools")); sys.path.insert(0, os.path.join(ROOT, "a7", "build"))
import tex2text
with contextlib.redirect_stdout(io.StringIO()):
    import scope_calendar as SC
specs = {}
for f in glob.glob(f"{ROOT}/a7/build/u*/l*.py"):
    s = importlib.util.spec_from_file_location("m", f); m = importlib.util.module_from_spec(s); s.loader.exec_module(m); specs[m.L["code"]] = m.L
stem = lambda code, u: (f"{u}.T{code[-1]}" if code.startswith("T-") else re.split(r"[+–]", code)[0])

wb = load_workbook(XLSX, data_only=True)
wbf = load_workbook(XLSX)                  # formulas
print("tabs:", [s.title + ("(hidden)" if s.sheet_state == "hidden" else "") for s in wb])

# ---------------- Scope & Sequence
ws = wb["Scope & Sequence"]; hdr = [c.value for c in ws[1]]; H = {h: i for i, h in enumerate(hdr)}
rows = list(ws.iter_rows(min_row=2))
if len(rows) != len(cal): bad("scope rows", len(rows), "calendar", len(cal))
prev = None; nlinks = 0; built_rows = 0
for i, (cells, c) in enumerate(zip(rows, cal), 2):
    v = {h: cells[j].value for h, j in H.items()}
    dt = D(v["Date"])
    if dt != d(c[0]): bad(f"r{i} date")
    if dt.strftime("%a") != c[1][:3]: bad(f"r{i} weekday", dt, c[1])
    if prev and dt <= prev: bad(f"r{i} order")
    prev = dt
    if v["Day"] != c[1] or (v["Unit"] or "") != (int(c[2]) if c[2] else "") or v["Lesson"] != c[3]: bad(f"r{i} day/unit/code")
    want = re.sub(r"\*", "", re.sub(r"^\*{1,2}\w+\*{1,2}\s*", "", c[4])).strip()
    if v["What is taught"] != want: bad(f"r{i} taught", repr(v["What is taught"])[:60], repr(want)[:60])
    if (v["Benchmark"] or "") != c[5] or (v["IXL skills — all required, code"] or "") != c[6]: bad(f"r{i} bm/ixl")
    if (D(v["IXL due"]) if v["IXL due"] else "") != (d(c[7]) if c[7] else ""): bad(f"r{i} due")
    code, u = c[3], int(c[2]) if c[2] else None
    exam = bool(re.fullmatch(r"[\d/]+\.X\d", code))
    # links
    for h in hdr[5:16]:
        cell = cells[H[h]]
        if cell.hyperlink:
            nlinks += 1
            kind, path = target(cell.hyperlink.target)
            if kind == "blob" and not exam and code != "spiral":
                name = path.rsplit("/", 1)[-1]; s = stem(code, u)
                want_name = {"Slides (PDF)": f"A7 {s}  Slides.pdf", "Slides (PowerPoint)": f"A7 {s}  Slides.pptx",
                             "Teacher Edition": f"A7 {s}  Teacher Edition.pdf", "Lesson Plan": f"A7 {s}  Lesson Plan.pdf",
                             "Question Bank": f"A7 {s}  Question Bank.pdf", "Bank key": f"A7 {s}  Question Bank Key.pdf",
                             "Bank – Additional": f"A7 {s}  Question Bank - Additional.pdf",
                             "Additional key": f"A7 {s}  Question Bank - Additional Key.pdf",
                             "Independent Set": f"A7 {s}  Independent Set.pdf", "Independent key": f"A7 {s}  Independent Set Key.pdf"}[h]
                if name != want_name: bad(f"r{i} {code} {h} ->", name)
                if not path.startswith(f"a7/packages/A7 Unit {u} - "): bad(f"r{i} wrong unit folder", path)
    if exam or code == "spiral":
        t1 = cells[H["Slides (PDF)"]].hyperlink; t2 = cells[H["Bank key"]].hyperlink
        if u in (3, 4):
            w1, w2 = (f"A7 {u}  Unit Assessment.pdf", f"A7 {u}  Unit Assessment Key.pdf") if exam else (f"A7 {u}  Unit Review.pdf", f"A7 {u}  Unit Review Key.pdf")
            if not (t1 and urllib.parse.unquote(t1.target).endswith(w1)) or not (t2 and urllib.parse.unquote(t2.target).endswith(w2)): bad(f"r{i} {code} paper/key links")
    lesson_like = not exam and code not in ("spiral", "flex") and not code.startswith("PM")
    if u in (3, 4) and lesson_like:
        built_rows += 1
        missing = [h for h in hdr[5:16] if not cells[H[h]].hyperlink]
        if missing: bad(f"r{i} {code} missing links", missing)
        L = specs[stem(code, u)]
        if v["Learning target"] != L["target"] or v["Essential question"] != L["essential"]: bad(f"r{i} target/EQ")
        if v["Vocabulary"] != ", ".join(t for t, _ in L["vocab"]): bad(f"r{i} vocab")
        b9 = L["whiteboard"][8]
        if v["Closure question (board 9)"] != b9["qtext"] or v["Closure answer"] != b9["answer"]: bad(f"r{i} closure")
        mt = v["MTRs (tap for where they show up)"].split("\n")
        if [x.split(" ")[0] for x in mt] != [f"MA.K12.{k}" for k, _ in L["mtr"]]: bad(f"r{i} MTR codes")
        for x in mt:
            title = x.split(" ", 1)[1]
            if std_n and norm(title) not in std_n: bad(f"r{i} MTR title not in the standards:", title)
        if not cells[H["MTRs (tap for where they show up)"]].comment or not cells[H["Closure question (board 9)"]].comment: bad(f"r{i} comment missing")
        # slide 1 names the lesson
        pth = urllib.parse.unquote(cells[H["Slides (PDF)"]].hyperlink.target.split("/main/", 1)[1])
        head = pymupdf.open(f"{ROOT}/{pth}")[0].get_text().split("\n")[0].strip()
        lab = L.get("label") or f"Lesson {L['lesson_no']}"
        if head != f"GRADE 7 ACCELERATED  ·  UNIT {u}  ·  {lab.upper()}": bad(f"r{i} slide 1 reads", head)
    if u and u >= 5 and lesson_like:
        if v["Slides (PDF)"] != "not built yet" or any(cells[H[h]].hyperlink for h in hdr[5:16]): bad(f"r{i} unbuilt row")
    # benchmark wording
    for part in (v["Benchmark wording"] or "").split("\n\n"):
        if not part: continue
        code_b, text_b = part.split(" — ", 1)
        if code_b in SOT:
            if text_b != SOT[code_b]: bad(f"r{i} wording differs for", code_b)
        elif not (code_b.startswith("MA.912") or re.fullmatch(r"MA\.\d+\.[A-Z]+\.\d+", code_b)): bad(f"r{i} unknown benchmark", code_b)
    toks = [b.strip() for b in c[5].split("·") if b.strip()]
    got = [p.split(" — ")[0] for p in (v["Benchmark wording"] or "").split("\n\n") if p]
    exp = []
    for b in toks:
        m = re.fullmatch(r"(\d+)\.([A-Z]+)\.(\d+)\.(\d+)–(\d+)\.(\d+)", b)
        exp += [f"MA.{m[1]}.{m[2]}.{m[3]}.{k}" for k in range(int(m[4]), int(m[6]) + 1)] if m else ["MA." + b]
    if got != exp: bad(f"r{i} wording codes", got, exp)
    # Math Nation lessons: each line 'N.M Title' must be a lesson of the book, and match the calendar code
    mn = [x for x in (v["Math Nation lesson(s)"] or "").split("\n") if x]
    for x in mn:
        if x not in toc_txt: bad(f"r{i} MN line not in the book's contents:", x)
    nums = [t for x in mn for t in re.findall(r"^(\d+\.\d+)(?: / (\d+\.\d+))?", x)[0] if t]
    if lesson_like:
        if code.startswith("T-"):
            want = {"T-A1": ["14.1", "14.2"], "T-A2": ["14.3", "14.4"], "T-B1": ["16.1"], "T-B2": ["16.2"], "T-C1": ["17.1"], "T-C2": ["17.2", "17.3"]}[code]
        else:
            mm = re.match(r"(\d+)\.(\d+)(.*)$", code); U_, a, rest = int(mm[1]), int(mm[2]), mm[3]
            extra = [int(x) for x in re.findall(r"\d+", rest)]
            want = [a] + extra if "–" not in rest else list(range(a, extra[-1] + 1))
            want = [f"{U_}.{k}" for k in want]
        want = [w for w in want if re.search(r"(?m)(^|[/·] )" + re.escape(w) + r"( |$)", toc_txt)]
        if nums != want: bad(f"r{i} {code} book lessons", nums, want)
print(f"scope: {len(rows)} rows, {nlinks} links, {built_rows} built lesson rows checked field by field")

# ---------------- Source of Truth statements: complete and verbatim against the official stext
# Two statements cannot be matched automatically and were checked by eye against the published pages
# on 21 Sep 2026: the PDF stext loses the superscripts in 8.AR.2.3 (x² = p, x³ = q), and a fraction
# from 7.NSO.1.1's clarification spills into 7.NSO.1.2.
SEEN = {"MA.8.AR.2.3", "MA.7.NSO.1.2"}
if std:
    import unicodedata
    raw = unicodedata.normalize("NFKC", std)
    CODE = r"MA\s*\.\s*(?:K12|\d+)\s*\.\s*[A-Z](?:\s*[A-Z])*\s*\.\s*\d+\s*\.\s*\d+(?!\s*\.\s*\d)"
    lines = []
    for l in raw.splitlines():
        if re.match(r"\s*=====.*page \d+ =====", l):
            lines.append("¶"); continue                  # a page break is a boundary
        lines.append(re.sub(CODE, " ", l))
    stext = re.sub(r"\s+-\s+", "-", " ".join(lines))
    out, pos = [], []
    for i, ch in enumerate(stext):
        c = ch.lower()
        if c.isascii() and c.isalnum():
            out.append(c); pos.append(i)
    T = "".join(out)
    path = os.path.join(REF, "Florida BEST Grade 8 - Source of Truth.md")
    sotn = unicodedata.normalize("NFKC", open(path, encoding="utf-8").read())
    SOTN = {m.group(1): re.sub(r"\*", "", m.group(2)).strip() for m in re.finditer(r"^\*\*(MA\.[\w.]+)\*\* — (.+)$", sotn, re.M)}
    good = 0; report = []
    for k, t in SOTN.items():
        n = "".join(c for c in t.lower() if c.isascii() and c.isalnum())
        starts = [m.start() for m in re.finditer(re.escape(n), T)]
        ok = False; ctx = None
        for j in starts:
            a, b = pos[j], pos[j + len(n) - 1] + 1
            before, after = stext[max(0, a - 70):a], stext[b:b + 60]
            bnd_before = re.search(r"([.?:¶)]|^)\s*$", before) or re.search(r"MA\.\d+\.[A-Z]+\.\d+\s+[A-Z][^.]*\.\s*$", before)
            bnd_after = re.match(r"\s*(\.|¶|Example|Benchmark|Clarification|MA\.|$)", after)
            ctx = (before, after)
            if bnd_before and bnd_after:
                ok = True; break
        if ok: good += 1
        else: report.append((k, t, ctx, len(starts)))
    print(f"Source of Truth: {good} of {len(SOTN)} statements complete and verbatim; {len(report)} checked by eye: {[k for k, *_ in report]}")
    for k, t, ctx, n in report:
        if k not in SEEN: bad("Source of Truth statement not complete/verbatim:", k, t)

# ---------------- Links (hidden) mirror the Scope links; HYPERLINK limit 255
wl = wb["Links"]; long_ = 0
for i, (cells, lrow) in enumerate(zip(rows, list(wl.iter_rows(min_row=2))), 2):
    for j, h in enumerate(hdr[5:16], 1):
        a = cells[H[h]].hyperlink.target if cells[H[h]].hyperlink else None
        b = lrow[j].value
        if a != b: bad(f"Links r{i} {h} mismatch")
        if b and len(b) > 255: long_ += 1
if long_: bad(f"{long_} URLs longer than 255 characters (HYPERLINK limit)")

# ---------------- IXL tracker
wi = wb["IXL tracker"]; irows = [r for r in wi.iter_rows(min_row=5) if r[0].value is not None]
if len(irows) != len(due): bad("IXL rows", len(irows), len(due))
for k, (r, x) in enumerate(zip(irows, due), 1):
    if r[0].value != k or D(r[1].value) != d(x[0]) or r[2].value != x[1] or r[3].value != x[2] or D(r[6].value) != d(x[3]) or r[5].value != 67: bad(f"IXL #{k} fields")
    if r[4].value != ", ".join(re.findall(r"\(([A-Z0-9]{3})\)", x[2])): bad(f"IXL #{k} codes", r[4].value)
    if r[9].value != "upcoming" and d(x[0]) > datetime.date.today(): bad(f"IXL #{k} status", r[9].value)
    loc = r[2].hyperlink.location if r[2].hyperlink else None
    m = re.fullmatch(r"'Scope & Sequence'!D(\d+)", loc or "")
    if not m or D(ws.cell(row=int(m[1]), column=1).value) != d(x[0]) or not x[1].startswith(ws.cell(row=int(m[1]), column=4).value): bad(f"IXL #{k} internal link", loc)
print(f"IXL tracker: {len(irows)} assignments")

# ---------------- Unit tests
wt = wb["Unit tests"]
from lib import unitbuild
exam_dates = collections.defaultdict(list)
for c in cal:
    if re.fullmatch(r"[\d/]+\.X\d", c[3]): exam_dates[c[3].split(".X")[0]].append(d(c[0]))
cal_rows = {r[0].value: r for r in wt.iter_rows(min_row=4, max_row=25) if isinstance(r[0].value, str) and r[0].value.startswith("Unit ")}
for k, ds in exam_dates.items():
    r = cal_rows.get(f"Unit {k}")
    if not r or [D(r[1].value), D(r[2].value)] != ds: bad("test dates", k, ds)
units = {}
for f in glob.glob(f"{ROOT}/a7/build/u*/unit.py"):
    s = importlib.util.spec_from_file_location("m", f); m = importlib.util.module_from_spec(s); s.loader.exec_module(m)
    U = [v for v in vars(m).values() if isinstance(v, dict) and "assessment" in v][0]; units[U["unit"]] = U
for u, U in sorted(units.items()):
    if cal_rows[f"Unit {u}"][3].value != U["assessment"]["total"]: bad(f"Unit {u} points", cal_rows[f'Unit {u}'][3].value)
    # the item table under the Unit u heading
    start = next(r[0].row for r in wt.iter_rows() if r[0].value == f"Unit {u} test — question map")
    items = []
    for r in wt.iter_rows(min_row=start + 3):
        if not isinstance(r[0].value, int): break
        items.append(r)
    n = 0; exp = []
    for sec in U["assessment"]["sections"]:
        for it in sec["items"]:
            if it.get("heading"): continue
            n += 1; pts = len(it["parts"]) if it.get("parts") else 1
            exp.append((n, pts, sec["benchmark"], bool(it.get("transfer")), "Multiple choice" if it.get("choices") else ("Parts" if it.get("parts") else "Written")))
    got = [(r[0].value, r[2].value, r[4].value, r[6].value == "yes", r[5].value) for r in items]
    if got != exp: bad(f"Unit {u} item map differs")
    summ = {wt.cell(row=start + 2 + k, column=9).value: wt.cell(row=start + 2 + k, column=10).value for k in range(1, 6)}
    for bm, pts in U["assessment"]["tracker"].items():
        if summ.get(bm) != pts: bad(f"Unit {u} {bm} points", summ.get(bm), pts)
    print(f"Unit {u} test: {len(items)} questions, {sum(x[1] for x in exp)} points; by benchmark {U['assessment']['tracker']} matches the sheet")

# ---------------- Standards
wsd = wb["Standards"]; srows = list(wsd.iter_rows(min_row=2))
codes_seen = [r[0].value for r in srows]
if len(set(codes_seen)) != len(codes_seen): bad("duplicate standards rows")
if set(SOT) - set(codes_seen): bad("SoT benchmark missing from Standards", set(SOT) - set(codes_seen))
test_pts = {bm: p for U in units.values() for bm, p in U["assessment"]["tracker"].items()}
for r in srows:
    code = r[0].value
    if code in SOT and r[2].value != SOT[code]: bad("Standards wording", code)
    want_pts = test_pts.get(code, "")
    if (r[9].value if r[9].value is not None else "") != want_pts: bad("Standards points", code, r[9].value, want_pts)
    short = code[3:]
    days = [c for c in cal if not c[3].startswith("PM") and short in [b.strip() for b in c[5].split("·")]]
    if days:
        if D(r[6].value) != min(d(c[0]) for c in days) or D(r[7].value) != max(d(c[0]) for c in days): bad("Standards first/last", code)
        for c in days:
            if f"{c[3]} (" not in (r[5].value or ""): bad("Standards taught-on misses", code, c[3])
print(f"Standards: {len(srows)} benchmarks")

# ---------------- Month calendar: every class day once, in its weekday column; holidays marked
wm = wb["Month calendar"]; seen = collections.Counter(); month = None
for row in wm.iter_rows():
    a = row[0].value
    if isinstance(a, str) and re.fullmatch(r"[A-Z][a-z]+ \d{4}", a):
        month = datetime.datetime.strptime(a, "%B %Y").date(); continue
    for j, cell in enumerate(row[:5]):
        if not isinstance(cell.value, str) or not re.match(r"^\d+", cell.value): continue
        dn = int(re.match(r"^(\d+)", cell.value)[1]); day = month.replace(day=dn)
        if day.weekday() != j: bad("calendar weekday", day, j)
        body = cell.value.split("\n", 1)[1] if "\n" in cell.value else ""
        seen[day] += 1
        c = next((x for x in cal if d(x[0]) == day), None)
        if c:
            if f"   {c[3]}\n" not in cell.value + "\n" and not cell.value.startswith(f"{dn}   {c[3]}"): bad("calendar code", day, cell.value[:30])
        elif day in SC.hol:
            if body != "No school": bad("holiday not marked", day, body)
        elif day >= datetime.date(2027, 5, 3) and day in SC.days:
            if body != "PM3 window": bad("PM3 day", day, body)
        elif day > SC.end:
            bad("labelled a day after school ends", day)
        elif day < SC.start:
            if body != "Units 1–2": bad("pre-plan day", day)
        else:
            bad("calendar day with no source", day, cell.value)
for c in cal:
    if seen[d(c[0])] != 1: bad("class day not once in the grid", c[0], seen[d(c[0])])
wkdays = [SC.start + datetime.timedelta(k) for k in range((SC.end - SC.start).days + 1)]
for day in wkdays:
    if day.weekday() < 5 and seen[day] != 1: bad("weekday missing from grid", day)
print(f"month calendar: {sum(seen.values())} weekday cells")

# ---------------- Vocabulary
wv = wb["Vocabulary"]; vrows = [r for r in wv.iter_rows(min_row=2) if r[3].value]
exp = [(t, tex2text.text(df)) for code in sorted(specs, key=lambda c: min(d(x[0]) for x in cal if stem(x[3], int(x[2]) if x[2] else 0) == c)) for t, df in specs[code]["vocab"]]
if [(r[3].value, r[4].value) for r in vrows] != exp: bad("vocabulary differs")
print(f"vocabulary: {len(vrows)} terms")

# ---------------- Units
wu = wb["Units"]; counts = collections.Counter(int(c[2]) for c in cal if c[2])
for r in wu.iter_rows(min_row=2, max_row=18):
    u = r[0].value
    if counts.get(u) and r[4].value != counts[u]: bad("Units periods", u, r[4].value, counts[u])
    for x in r[6:]:
        if x.hyperlink: target(x.hyperlink.target)
for r in wb["Units 1–2 decks"].iter_rows(min_row=2):
    for x in r[2:]:
        if x.hyperlink: target(x.hyperlink.target)
for r in wb["Unit tests"].iter_rows():
    for x in r:
        if x.hyperlink and x.hyperlink.target: target(x.hyperlink.target)

# ---------------- whole workbook: no LaTeX left, one font, no error strings
for s in wbf:
    for row in s.iter_rows():
        for x in row:
            vals = [x.value] if isinstance(x.value, str) and not x.value.startswith("=") else []
            if x.comment: vals.append(x.comment.text)
            for t in vals:
                if re.search(r"\$|\\[a-z]+|\{|\}", t): bad("LaTeX left in", s.title, x.coordinate, t[:60])
            if x.value is not None and x.font.name != "Arial": bad("font", s.title, x.coordinate, x.font.name)
for s in wb:
    for row in s.iter_rows():
        for x in row:
            if isinstance(x.value, str) and re.fullmatch(r"#(REF!|VALUE!|NAME\?|N/A|DIV/0!|NUM!)", x.value): bad("error value", s.title, x.coordinate)
print("\n".join(P) if P else "no problems"); print(len(P), "problems")
sys.exit(1 if P else 0)
