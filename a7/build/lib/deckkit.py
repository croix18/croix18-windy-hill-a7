"""pptx builder matching the shipped A7 deck geometry (13.33 x 7.5 in, Century Schoolbook,
INK/VOCAB/RED/GRAY, double rule under the title, footer rule + running title + page number).
No speaker notes in the deck (ruling 12): every slide's minutes and teaching note go to a
side-car JSON beside the .pptx, which the teacher's edition reads at build time.
"""
import json, os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from . import mathimg

INK, VOCAB, RED, GRAY, LT, FILL = "1A1A1A", "0B5394", "9E1B32", "6B6B6B", "D9D9D9", "F2F2F0"
FONT = "Century Schoolbook"
W, H = 13.3333, 7.5
LM, CW = 0.85, 11.6
FOOT_Y = 6.78


def _rgb(h):
    return RGBColor.from_string(h)


SUP = dict(zip("⁰¹²³⁴⁵⁶⁷⁸⁹⁻⁺ᵐⁿ", "0123456789-+mn"))


def _split_sup(text):
    """Unicode superscript digits have no glyph in the deck font (glyph check); emit them as
    real superscript runs instead. Returns [(piece, is_sup), ...]."""
    out = []; buf = ""; cur = False
    for ch in text:
        sup = ch in SUP
        if sup != cur and buf:
            out.append((buf, cur)); buf = ""
        cur = sup
        buf += SUP[ch] if sup else ch
    if buf:
        out.append((buf, cur))
    return out


class Deck:
    def __init__(self, course, unit, lesson_label, title, footer):
        self.p = Presentation()
        self.p.slide_width = Inches(W); self.p.slide_height = Inches(H)
        self.blank = self.p.slide_layouts[6]
        self.footer = footer
        self.course, self.unit, self.lesson_label, self.title = course, unit, lesson_label, title
        self.side = []   # side-car records
        self.s = None
        self.cursor = 0.0
        self._n = 0

    # ---------- primitives ----------
    def _text(self, x, y, w, h, text, size=23, bold=False, italic=False, color=INK, align="left",
              anchor="top", runs=None, wrap=True, foot=False):
        if not foot and y + h > FOOT_Y + 0.02:
            raise RuntimeError(f"text box runs into the footer (bottom {y + h:.2f} in): {(text or (runs and runs[0][0]) or '')[:50]}")
        if x < 0 or x + w > W + 0.01:
            raise RuntimeError(f"text box off the slide horizontally: {(text or '')[:50]}")
        tb = self.s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame; tf.word_wrap = wrap
        tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = Inches(0.02)
        tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE}[anchor]
        para = tf.paragraphs[0]
        para.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
        for (t, sz, b, i, c) in (runs or [(text, size, bold, italic, color)]):
            for piece, sup in _split_sup(t):
                r = para.add_run(); r.text = piece
                r.font.name = FONT; r.font.size = Pt(sz); r.font.bold = b; r.font.italic = i
                r.font.color.rgb = _rgb(c)
                if sup:
                    r.font._element.set("baseline", "30000")
        return tb

    def _line(self, x, y, w, weight=1.5, color=INK):
        ln = self.s.shapes.add_connector(1, Inches(x), Inches(y), Inches(x + w), Inches(y))
        ln.line.color.rgb = _rgb(color); ln.line.width = Pt(weight)
        return ln

    def _rect(self, x, y, w, h, fill=FILL, line=None):
        r = self.s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        r.fill.solid(); r.fill.fore_color.rgb = _rgb(fill)
        if line:
            r.line.color.rgb = _rgb(line); r.line.width = Pt(0.75)
        else:
            r.line.fill.background()
        r.shadow.inherit = False
        return r

    def _foot(self):
        self._line(LM, FOOT_Y, CW, 0.75, GRAY)
        self._text(LM, FOOT_Y + 0.08, 10.6, 0.3, self.footer, 11, italic=True, color=GRAY, foot=True)
        self._text(11.75, FOOT_Y + 0.08, 0.7, 0.3, str(self._n), 11, color=GRAY, align="right", foot=True)

    def _new(self, title, sub, minutes, note, kind):
        self.s = self.p.slides.add_slide(self.blank)
        self._n += 1
        self.side.append({"n": self._n, "title": title, "sub": sub, "min": minutes, "note": note, "kind": kind})
        self._text(LM, 0.4, CW, 0.7, title, 36 if len(title) < 30 else 34, bold=True)
        self._line(LM, 1.2, CW, 1.5); self._line(LM, 1.29, CW, 0.75)
        if sub:
            self._text(LM, 1.38, CW, 0.34, sub, 17, italic=True, color=GRAY)
        self.cursor = 1.9
        self._foot()
        return self.s

    # ---------- slide types ----------
    def title_slide(self, benchmark, target, yesterday, today, minutes=1, note=""):
        self.s = self.p.slides.add_slide(self.blank); self._n += 1
        self.side.append({"n": self._n, "title": self.title, "sub": "", "min": minutes, "note": note, "kind": "title"})
        eyebrow = f"{self.course}  ·  UNIT {self.unit}  ·  {self.lesson_label}".upper()
        self._text(LM, 1.85, CW, 0.38, eyebrow, 15, color=GRAY, align="center")
        self._text(LM, 2.24, CW, 1.1, self.title, 34, bold=True, align="center", anchor="middle")
        self._line(3.6, 3.56, 6.1, 1.5); self._line(3.6, 3.65, 6.1, 0.75)
        self._text(LM, 3.83, CW, 0.34, benchmark, 16, bold=True, align="center")
        tl = max(1, -(-len(target) // 84))
        self._text(LM, 4.2, CW, 0.4 * tl, target, 19, italic=True, align="center")
        by = 4.8 + 0.38 * (tl - 1)
        self._rect(2.6, by, 8.1, 1.15, FILL, "BFBFBF")
        self._text(2.85, by + 0.17, 7.6, 0.4, yesterday, 17, italic=True, color=GRAY, align="center")
        self._text(2.85, by + 0.59, 7.6, 0.4, today, 17, bold=True, align="center")
        self._foot()

    def section(self, title, sub="", minutes=0, note="", kind="content"):
        return self._new(title, sub, minutes, note, kind)

    def head(self, text, numeral=None):
        """A bold sub-heading with a thin rule (Notes I. / II.)."""
        label = (f"{numeral}.  " if numeral else "") + text
        self._text(LM, self.cursor, CW, 0.44, label, 25, bold=True)
        self._line(LM, self.cursor + 0.48, CW, 0.75)
        self.cursor += 0.66

    def items(self, rows, size=23, panel=False, letters=True, x=None, w=None, gap=0.1, start=0):
        """Lettered rows. Each row: plain string, or (term, rest) for a vocab row (term blue bold),
        or ("**bold**", ...) — a row beginning with ** is set bold. Returns bottom y."""
        x = LM + 0.11 if x is None else x
        w = CW - 0.22 if w is None else w
        y = self.cursor
        for i, row in enumerate(rows):
            if isinstance(row, tuple):
                term, rest = row
                runs = [(term, size + 2, True, False, VOCAB), ("   —   " + rest, size, False, False, INK)]
                text = term + "   —   " + rest
            else:
                bold = row.startswith("**")
                text = row.strip("*")
                runs = [(text, size, bold, False, INK)]
            if "$" in text:
                xx = x + (0.77 if letters else 0.12)
                if letters:
                    self._text(x + 0.12, y + 0.07, 0.6, 0.42, f"{chr(65 + start + i)}.", size, color=INK)
                tw, h = self._mixed(text, xx, y + 0.05, "slidemid", size, INK, bold)
                y += h + 0.18 + gap
                continue
            lines = max(1, int(len(text) * (size / 23) / 82) + 1)
            h = 0.52 * lines
            if panel:
                self._rect(x - 0.03, y - 0.02, w, h + 0.12)
            if letters:
                self._text(x + 0.12, y + 0.07, 0.6, 0.42, f"{chr(65 + start + i)}.", size, color=INK)
                self._text(x + 0.77, y + 0.07, w - 0.95, h, "", size, runs=runs)
            else:
                self._text(x + 0.12, y + 0.07, w - 0.3, h, "", size, runs=runs)
            y += h + 0.18 + gap
        self.cursor = y
        return y

    def numbered(self, rows, size=21, x=1.05, w=11.4, gap=0.2):
        y = self.cursor
        for i, text in enumerate(rows):
            cpl = int((w - 0.65) * 72 / (size * 0.50))
            lines = max(1, -(-len(text) // cpl))
            h = 0.39 * lines + 0.06
            self._text(x, y, 0.6, 0.39, f"{i + 1}.", size, bold=True, color=VOCAB)
            self._text(x + 0.65, y, w - 0.65, h, text, size)
            y += h + gap
        self.cursor = y

    def text(self, text, size=23, bold=False, italic=False, color=INK, align="left", h=None, x=None, w=None):
        x = LM if x is None else x; w = CW if w is None else w
        cpl = int(w * 72 / (size * 0.50))
        lines = max(1, -(-len(text) // cpl))
        h = h or (0.45 * lines * (size / 23))
        self._text(x, self.cursor, w, h, text, size, bold, italic, color, align)
        self.cursor += h + 0.12
        return self.cursor

    def table(self, widths, rows, size=15, header=True, x=None, row_h=0.42):
        """Simple table: widths in inches; rows of strings. Header row shaded."""
        x = LM + (CW - sum(widths)) / 2 if x is None else x
        shp = self.s.shapes.add_table(len(rows), len(widths), Inches(x), Inches(self.cursor),
                                      Inches(sum(widths)), Inches(row_h * len(rows)))
        tbl = shp.table
        tbl.first_row = False; tbl.horz_banding = False
        for ci, wv in enumerate(widths):
            tbl.columns[ci].width = Inches(wv)
        for ri, row in enumerate(rows):
            tbl.rows[ri].height = Inches(row_h)
            for ci, val in enumerate(row):
                cell = tbl.cell(ri, ci)
                cell.fill.solid(); cell.fill.fore_color.rgb = _rgb(LT if (header and ri == 0) else "FFFFFF")
                cell.margin_left = cell.margin_right = Inches(0.08)
                cell.margin_top = cell.margin_bottom = Inches(0.03)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf = cell.text_frame; tf.word_wrap = True
                para = tf.paragraphs[0]; para.alignment = PP_ALIGN.CENTER
                for piece, sup in _split_sup(val):
                    r = para.add_run(); r.text = piece
                    r.font.name = FONT; r.font.size = Pt(size); r.font.bold = (header and ri == 0)
                    r.font.color.rgb = _rgb(INK)
                    if sup:
                        r.font._element.set("baseline", "30000")
        # borders: python-pptx has no API; set via XML
        from pptx.oxml.ns import qn
        from lxml import etree
        for ri in range(len(rows)):
            for ci in range(len(widths)):
                tc = tbl.cell(ri, ci)._tc
                tcPr = tc.get_or_add_tcPr()
                for edge in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
                    ln = etree.SubElement(tcPr, qn(edge), w="9525", cap="flat", cmpd="sng", algn="ctr")
                    sf = etree.SubElement(ln, qn("a:solidFill"))
                    etree.SubElement(sf, qn("a:srgbClr"), val=INK)
                    etree.SubElement(ln, qn("a:prstDash"), val="solid")
        self.cursor += row_h * len(rows) + 0.25
        return tbl

    def math(self, latex, surface="slidebig", align="center", x=None, y=None, color=INK, gap=0.25):
        path, w, h = mathimg.m(latex, surface, color)
        yy = self.cursor if y is None else y
        if yy + h > FOOT_Y:
            raise RuntimeError(f"figure runs into the footer: {latex[:40]}")
        xx = x if x is not None else (LM + (CW - w) / 2 if align == "center" else LM + 0.3)
        self.s.shapes.add_picture(path, Inches(xx), Inches(yy), Inches(w), Inches(h))
        if y is None:
            self.cursor = yy + h + gap
        return w, h

    def _mixed(self, text, x, y, surface="slidemid", size=26, color=INK, bold=False, align="left", width=None):
        """Lay out one line mixing text and $latex$ pieces, images vertically centred on the text.
        Returns (total_width, line_height). Text width is estimated from character count and the
        pieces are placed left to right; align='center' centres the whole line inside [x, x+width]."""
        import re
        pieces = []
        total = 0.0; maxh = 0.45
        for part in re.split(r"(\$[^$]+\$)", text):
            if not part:
                continue
            if part.startswith("$"):
                path, w, h = mathimg.m(part[1:-1], surface, color)
                pieces.append(("img", path, w, h)); total += w; maxh = max(maxh, h)
            else:
                w = len(part) * size / 72 * (0.58 if bold else 0.52) + 0.08
                pieces.append(("txt", part, w, 0.5)); total += w
        if align == "center":
            x = x + ((width if width else CW) - total) / 2
        if x + total > LM + CW + 0.05:
            raise RuntimeError(f"line runs off the slide ({x + total:.2f} in): {text[:60]}")
        for kind, val, w, h in pieces:
            if kind == "img":
                self.s.shapes.add_picture(val, Inches(x), Inches(y + (maxh - h) / 2), Inches(w), Inches(h))
            else:
                self._text(x, y + (maxh - 0.5) / 2, w + 0.15, 0.5, val, size, bold=bold, color=color, anchor="middle", wrap=False)
            x += w
        return total, maxh

    def measure(self, text, surface="slidemid", size=26, bold=False):
        """Width in inches a mixed line would take, without drawing it."""
        import re
        total = 0.0
        for part in re.split(r"(\$[^$]+\$)", text):
            if not part:
                continue
            if part.startswith("$"):
                _, w, h = mathimg.m(part[1:-1], surface, INK)
                total += w
            else:
                total += len(part) * size / 72 * (0.58 if bold else 0.52) + 0.08
        return total

    def math_row(self, parts, surface="slidemid", y=None, gap=0.35, size=26, color=INK, bold=False, align="center", x=None):
        """A row mixing text and $latex$ pieces, vertically aligned; centered unless align='left'."""
        yy = self.cursor if y is None else y
        total, maxh = self._mixed(parts, LM if x is None else x, yy, surface, size, color, bold, align=align)
        if y is None:
            self.cursor = yy + maxh + gap
        return maxh

    def answer_line(self, text, y=5.2):
        self._text(LM, y, CW, 0.6, "Answer:   " + text, 32, bold=True, color=RED, align="center")

    def answer_math(self, latex, y=5.1):
        path, w, h = mathimg.m(latex, "slidebig", RED)
        lw = 2.1
        x = LM + (CW - w - lw) / 2
        self._text(x, y + (h - 0.6) / 2, lw, 0.6, "Answer:", 32, bold=True, color=RED, anchor="middle")
        self.s.shapes.add_picture(path, Inches(x + lw), Inches(y), Inches(w), Inches(h))

    def choices(self, opts, size=24, correct=None, two_col=True):
        """A–D options. On an answer slide pass correct=index to colour it red."""
        y = self.cursor
        n = len(opts)
        if two_col and n == 4:
            cols = [(LM + 1.2, 5.0), (LM + 6.4, 5.0)]
            for i, o in enumerate(opts):
                cx, cw_ = cols[i % 2]
                row = i // 2
                col = RED if correct == i else INK
                self._text(cx, y + row * 0.75, cw_, 0.6, "", size, runs=[(f"{chr(65 + i)}.   ", size, True, False, col), (o, size, correct == i, False, col)])
            self.cursor = y + 2 * 0.75 + 0.2
        else:
            for i, o in enumerate(opts):
                col = RED if correct == i else INK
                self._text(LM + 1.2, y + i * 0.62, CW - 1.5, 0.55, "", size, runs=[(f"{chr(65 + i)}.   ", size, True, False, col), (o, size, correct == i, False, col)])
            self.cursor = y + n * 0.62 + 0.2

    def ixl(self, skills, minutes=5):
        self.section("IXL", "Last five minutes.", minutes, "IXL: whatever is not finished is tonight's practice; SmartScore 67.", "ixl")
        rows = ["Open IXL and start today's skills.",
                "Work on paper where the question needs work. The answer box does not show it.",
                "Stop at a SmartScore of 67 on each skill. Whatever you do not finish is tonight's practice."]
        self.cursor = 2.0
        self.numbered(rows, gap=0.1)
        self.cursor += 0.05
        self._text(LM + 0.2, self.cursor, CW - 0.4, 0.4, "Today's skills", 21, bold=True)
        self.cursor += 0.45
        for s in skills:
            self._text(LM + 0.6, self.cursor, CW - 1.0, 0.36, s, 19, color=INK)
            self.cursor += 0.36

    # ---------- finish ----------
    def save(self, path):
        self.p.save(path)
        with open(path[:-5] + ".notes.json", "w") as f:
            json.dump({"deck": os.path.basename(path), "slides": self.side}, f, indent=1)
        return path
